"""Build reproducible synthetic training datasets for the TCI prototype.

The source workbook is explicitly synthetic. Outputs must remain labelled
synthetic_training_only and may not be presented as MAPID, KAI, survey, or
property observations.
"""

from __future__ import annotations

import csv
import json
import math
import random
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


SEED = 20260911
ROOT = Path(__file__).resolve().parents[2]
RAW_WORKBOOK = ROOT / "ml" / "shared" / "raw" / "Dummy_Dataset_Volume_Penumpang_KRL_Jabodetabek.xlsx"
SYNTHETIC_STATUS = "synthetic_training_only"


def norm(values: list[float]) -> list[float]:
    low, high = min(values), max(values)
    if math.isclose(low, high):
        return [0.5 for _ in values]
    return [(value - low) / (high - low) for value in values]


def clamp(value: float, low: float, high: float) -> float:
    return max(low, min(high, value))


def slugify_station(name: str) -> str:
    replacements = {
        " ": "",
        "-": "",
        ".": "",
        "'": "",
    }
    slug = name.lower()
    for source, target in replacements.items():
        slug = slug.replace(source, target)
    aliases = {
        "universitaspancasila": "univpancasila",
        "universitasindonesia": "univindonesia",
        "durenkalibata": "durenkalibata",
        "pasarminggubaru": "pasarminggubaru",
        "tanjungpriok": "tanjungpriok",
        "tanahabang": "tanahabang",
        "tanahtinggi": "tanahtinggi",
        "kampungbandan": "kampungbandan",
        "pasarsenen": "pasarsenen",
        # Dataset and frontend use different public names for these same stations.
        "bnicity": "sudirmanbaru",
        "metlandtelagamurni": "telagamurni",
    }
    return aliases.get(slug, slug)


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        raise ValueError(f"No rows for {path}")
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def number(value: float, digits: int = 4) -> float:
    return round(value, digits)


def load_source() -> tuple[list[dict], list[dict]]:
    if not RAW_WORKBOOK.exists():
        raise FileNotFoundError(
            f"Source workbook not found: {RAW_WORKBOOK}. Copy the supplied workbook first."
        )
    workbook = load_workbook(RAW_WORKBOOK, data_only=True)
    master_sheet = workbook["Master_Stasiun"]
    monthly_sheet = workbook["Volume_Bulanan_2023-2025"]

    master_headers = [master_sheet.cell(2, column).value for column in range(1, 14)]
    monthly_headers = [monthly_sheet.cell(2, column).value for column in range(1, 9)]
    master = [
        dict(zip(master_headers, row))
        for row in master_sheet.iter_rows(min_row=3, values_only=True)
        if row[0]
    ]
    monthly = [
        dict(zip(monthly_headers, row))
        for row in monthly_sheet.iter_rows(min_row=3, values_only=True)
        if row[0]
    ]
    return master, monthly


def main() -> None:
    random.seed(SEED)
    master, monthly = load_source()

    code_key = "Kode Internal Stasiun"
    name_key = "Nama Stasiun"
    line_key = "Lintas Utama"
    type_key = "Tipe Stasiun (Ukuran/Peran)"
    lat_key = "Latitude"
    lng_key = "Longitude"
    city_key = "Kota / Kabupaten"
    connection_key = "Jumlah Lintas Terhubung"
    hub_key = "Status Hub Transit"
    distance_key = "Jarak ke Monas (km)"
    annual_key = "Volume Penumpang Tahunan 2025 (orang)"
    daily_key = "Rata-rata Volume Harian 2025 (orang/hari)"
    yoy_key = "Pertumbuhan YoY 2024→2025 (%)"
    period_key = "Periode (YYYY-MM)"
    monthly_volume_key = "Volume Penumpang Bulanan (orang)"
    monthly_daily_key = "Rata-rata Volume Harian (orang/hari)"

    annual_norm = dict(zip([row[code_key] for row in master], norm([float(row[annual_key]) for row in master])))
    growth_norm = dict(zip([row[code_key] for row in master], norm([float(row[yoy_key]) for row in master])))
    distance_norm = dict(zip([row[code_key] for row in master], norm([float(row[distance_key]) for row in master])))

    master_by_code = {row[code_key]: row for row in master}
    crosswalk_rows = []
    station_context = {}
    for row in master:
        code = row[code_key]
        hub = 1 if row[hub_key] == "Ya" else 0
        connections = int(row[connection_key])
        centrality = 1 - distance_norm[code]
        activity = annual_norm[code]
        growth = growth_norm[code]
        accessibility = clamp(
            0.18 + 0.28 * hub + 0.24 * min(connections / 3, 1) + 0.20 * centrality + random.gauss(0, 0.035),
            0.05,
            0.98,
        )
        business_base = max(5, round(12 + 125 * activity + 35 * centrality + 18 * hub + random.gauss(0, 7)))
        land_availability = clamp(
            0.70 - 0.35 * centrality + 0.12 * (1 - activity) + random.gauss(0, 0.06),
            0.08,
            0.95,
        )
        office_index = clamp(0.12 + 0.62 * centrality + 0.20 * hub + random.gauss(0, 0.10), 0.03, 0.98)
        residential_index = clamp(0.35 + 0.42 * (1 - centrality) + 0.10 * (1 - hub) + random.gauss(0, 0.10), 0.05, 0.98)
        education_index = clamp(0.12 + 0.42 * (1 - centrality) + random.gauss(0, 0.13), 0.02, 0.95)
        station_id = slugify_station(row[name_key])
        station_context[code] = {
            "station_id": station_id,
            "hub": hub,
            "connections": connections,
            "centrality": centrality,
            "activity": activity,
            "growth": growth,
            "accessibility": accessibility,
            "business_base": business_base,
            "land_availability": land_availability,
            "office_index": office_index,
            "residential_index": residential_index,
            "education_index": education_index,
        }
        crosswalk_rows.append(
            {
                "station_code": code,
                "station_id": station_id,
                "station_name": row[name_key],
                "line": row[line_key],
                "latitude": row[lat_key],
                "longitude": row[lng_key],
                "source_status": SYNTHETIC_STATUS,
            }
        )

    forecast_rows = []
    by_station_period = defaultdict(dict)
    for row in monthly:
        code = row[code_key]
        period = row[period_key]
        by_station_period[code][period] = float(row[monthly_volume_key])
        forecast_rows.append(
            {
                "station_code": code,
                "station_id": station_context[code]["station_id"],
                "station_name": row[name_key],
                "line": row[line_key],
                "year": int(row["Tahun"]),
                "month": int(row["Bulan (1-12)"]),
                "period": period,
                "passenger_volume_monthly": int(row[monthly_volume_key]),
                "passenger_volume_daily_avg": int(row[monthly_daily_key]),
                "source_status": SYNTHETIC_STATUS,
            }
        )

    cluster_rows = []
    for row in master:
        code = row[code_key]
        context = station_context[code]
        cluster_rows.append(
            {
                "station_code": code,
                "station_id": context["station_id"],
                "station_name": row[name_key],
                "line": row[line_key],
                "station_type": row[type_key],
                "city_regency": row[city_key],
                "latitude": row[lat_key],
                "longitude": row[lng_key],
                "connected_lines": context["connections"],
                "is_transit_hub": context["hub"],
                "distance_to_monas_km": row[distance_key],
                "annual_passengers_2025": row[annual_key],
                "daily_passengers_2025": row[daily_key],
                "yoy_growth_2025_pct": number(float(row[yoy_key]) * 100, 2),
                "source_status": SYNTHETIC_STATUS,
            }
        )

    regression_rows = []
    classification_rows = []
    latest_by_station = {}
    for source in sorted(monthly, key=lambda item: (item[code_key], item[period_key])):
        code = source[code_key]
        context = station_context[code]
        station = master_by_code[code]
        period = source[period_key]
        month = int(source["Bulan (1-12)"])
        passenger = float(source[monthly_volume_key])
        passenger_relative = passenger / max(by_station_period[code].values())
        season = math.sin((month - 1) * math.pi / 6)
        business_density = max(1, round(context["business_base"] * (0.72 + 0.26 * passenger_relative) + random.gauss(0, 4)))
        poi_density = max(1, round(business_density * (0.78 + 0.25 * context["centrality"]) + random.gauss(0, 5)))
        pedestrian_access = clamp(context["accessibility"] + random.gauss(0, 0.025), 0.03, 0.99)
        property_availability = clamp(context["land_availability"] + random.gauss(0, 0.035), 0.03, 0.98)
        intermodal = clamp(0.20 + 0.35 * context["hub"] + 0.23 * min(context["connections"] / 3, 1) + random.gauss(0, 0.03), 0.05, 0.98)
        rental_index = max(15, 32 + 85 * context["centrality"] + 35 * context["activity"] + random.gauss(0, 7))
        investment_score = clamp(
            100
            * (
                0.26 * passenger_relative
                + 0.12 * context["growth"]
                + 0.17 * pedestrian_access
                + 0.14 * min(business_density / 175, 1)
                + 0.15 * property_availability
                + 0.16 * intermodal
            )
            + random.gauss(0, 3.2),
            8,
            97,
        )
        priority = "High" if investment_score >= 70 else "Medium" if investment_score >= 45 else "Low"
        regression_row = {
            "record_id": f"REG-{code}-{period}",
            "station_code": code,
            "station_id": context["station_id"],
            "station_name": station[name_key],
            "line": station[line_key],
            "period": period,
            "passenger_volume_monthly": int(passenger),
            "passenger_growth_yoy_pct": number(float(station[yoy_key]) * 100 + random.gauss(0, 0.55), 2),
            "business_density_750m": business_density,
            "poi_density_750m": poi_density,
            "pedestrian_access_index": number(pedestrian_access * 100, 2),
            "land_property_availability_index": number(property_availability * 100, 2),
            "intermodal_connectivity_index": number(intermodal * 100, 2),
            "estimated_rental_index_rp_thousand_m2_month": number(rental_index, 2),
            "distance_to_cbd_km": station[distance_key],
            "investment_potential_score": number(investment_score, 2),
            "investment_priority": priority,
            "target_provenance": "synthetic_formula_with_seeded_noise",
            "source_status": SYNTHETIC_STATUS,
        }
        regression_rows.append(regression_row)
        latest_by_station[code] = regression_row

        radius = (500, 750, 1000)[(int(source["Tahun"]) + month + int(code.split("-")[1])) % 3]
        radius_factor = radius / 750
        footfall = max(90, int(passenger / 30 * (0.25 + 0.15 * pedestrian_access) * radius_factor + random.gauss(0, 75)))
        food_poi = max(0, round((0.28 * business_density + 12 * context["office_index"] + 7 * context["residential_index"]) * radius_factor + random.gauss(0, 3)))
        drink_poi = max(0, round((0.16 * business_density + 16 * context["office_index"] + 6 * context["education_index"]) * radius_factor + random.gauss(0, 3)))
        hobby_poi = max(0, round((0.08 * business_density + 11 * context["residential_index"] + 11 * context["education_index"]) * radius_factor + random.gauss(0, 2)))
        book_poi = max(0, round((0.025 * business_density + 15 * context["education_index"]) * radius_factor + random.gauss(0, 1.5)))
        other_poi = max(0, round((0.22 * business_density + 10 * context["centrality"]) * radius_factor + random.gauss(0, 4)))
        opportunity = {
            "Makanan": 0.34 * footfall / 1000 + 0.26 * context["office_index"] + 0.22 * context["residential_index"] - 0.018 * food_poi + 0.06 * season,
            "Minuman": 0.29 * footfall / 1000 + 0.37 * context["office_index"] + 0.20 * context["education_index"] - 0.020 * drink_poi + 0.05 * season,
            "Hobi": 0.16 * footfall / 1000 + 0.38 * context["residential_index"] + 0.36 * context["education_index"] - 0.022 * hobby_poi,
            "Toko Buku": 0.11 * footfall / 1000 + 0.58 * context["education_index"] + 0.12 * context["residential_index"] - 0.035 * book_poi,
            "Lainnya": 0.22 * footfall / 1000 + 0.30 * context["centrality"] + 0.12 * context["activity"] - 0.015 * other_poi,
        }
        selected_category = max(opportunity, key=opportunity.get)
        sorted_scores = sorted(opportunity.values(), reverse=True)
        confidence = clamp(0.50 + (sorted_scores[0] - sorted_scores[1]) * 0.32 + random.gauss(0, 0.03), 0.50, 0.92)
        classification_rows.append(
            {
                "record_id": f"CLS-{code}-{period}-{radius}",
                "station_code": code,
                "station_id": context["station_id"],
                "station_name": station[name_key],
                "period": period,
                "buffer_radius_m": radius,
                "estimated_daily_footfall": footfall,
                "passenger_volume_monthly": int(passenger),
                "food_poi_count": food_poi,
                "beverage_poi_count": drink_poi,
                "hobby_poi_count": hobby_poi,
                "bookstore_poi_count": book_poi,
                "other_retail_poi_count": other_poi,
                "office_activity_index": number(context["office_index"] * 100, 2),
                "residential_activity_index": number(context["residential_index"] * 100, 2),
                "education_activity_index": number(context["education_index"] * 100, 2),
                "pedestrian_access_index": number(pedestrian_access * 100, 2),
                "estimated_rental_index_rp_thousand_m2_month": number(rental_index, 2),
                "recommended_business_category": selected_category,
                "synthetic_recommendation_confidence": number(confidence, 3),
                "label_provenance": "synthetic_opportunity_formula_with_seeded_noise",
                "source_status": SYNTHETIC_STATUS,
            }
        )

    llm_context_rows = []
    llm_eval_rows = []
    category_by_station = {}
    for row in classification_rows:
        if row["period"] == "2025-12":
            category_by_station[row["station_code"]] = row
    for row in master:
        code = row[code_key]
        context = station_context[code]
        regression = latest_by_station[code]
        category = category_by_station[code]
        monthly_values = by_station_period[code]
        recent = [monthly_values[f"2025-{month:02d}"] for month in (10, 11, 12)]
        trend_pct = (recent[-1] / recent[0] - 1) * 100
        payload = {
            "record_id": f"LLMCTX-{code}-2025-12",
            "station_code": code,
            "station_id": context["station_id"],
            "station_name": row[name_key],
            "line": row[line_key],
            "as_of_period": "2025-12",
            "passenger_volume_recent_3m": [int(value) for value in recent],
            "passenger_trend_recent_3m_pct": number(trend_pct, 2),
            "investment_potential_score": regression["investment_potential_score"],
            "investment_priority": regression["investment_priority"],
            "typology_input": {
                "annual_passengers_2025": int(row[annual_key]),
                "connected_lines": context["connections"],
                "is_transit_hub": bool(context["hub"]),
            },
            "business_recommendation": {
                "category": category["recommended_business_category"],
                "confidence": category["synthetic_recommendation_confidence"],
                "buffer_radius_m": category["buffer_radius_m"],
            },
            "allowed_claims": [
                "Summarise supplied values only.",
                "State that all available inputs are synthetic prototypes.",
                "Do not give binding investment, permit, or legal advice.",
            ],
            "source_status": SYNTHETIC_STATUS,
        }
        llm_context_rows.append(payload)
        for intent, query in (
            ("station_summary", f"Ringkas potensi {row[name_key]} berdasarkan data yang tersedia."),
            ("business_category", f"Kategori usaha apa yang dapat dieksplorasi di sekitar {row[name_key]}?"),
            ("data_limitation", f"Seberapa dapat diandalkan insight untuk {row[name_key]}?"),
        ):
            llm_eval_rows.append(
                {
                    "case_id": f"LLMEVAL-{code}-{intent}",
                    "station_id": context["station_id"],
                    "intent": intent,
                    "user_query": query,
                    "required_fact_keys": json.dumps(
                        ["station_name", "source_status"]
                        + (["investment_potential_score"] if intent == "station_summary" else [])
                        + (["business_recommendation.category"] if intent == "business_category" else []),
                        ensure_ascii=False,
                    ),
                    "must_include_limitation": True,
                    "source_status": SYNTHETIC_STATUS,
                }
            )

    write_csv(ROOT / "ml" / "shared" / "processed" / "station_crosswalk.csv", crosswalk_rows)
    write_csv(ROOT / "ml" / "forecasting" / "data" / "passenger_volume_monthly.csv", forecast_rows)
    write_csv(ROOT / "ml" / "clustering" / "data" / "station_master.csv", cluster_rows)
    write_csv(ROOT / "ml" / "regression_scoring" / "data" / "station_month_features.csv", regression_rows)
    write_csv(ROOT / "ml" / "business_classification" / "data" / "location_opportunity_training.csv", classification_rows)

    llm_dir = ROOT / "ml" / "llm_insight" / "data"
    llm_dir.mkdir(parents=True, exist_ok=True)
    with (llm_dir / "station_insight_context.jsonl").open("w", encoding="utf-8") as file:
        for row in llm_context_rows:
            file.write(json.dumps(row, ensure_ascii=False) + "\n")
    write_csv(llm_dir / "smart_query_evaluation.csv", llm_eval_rows)

    manifest = {
        "generated_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "generator": "ml/scripts/generate_synthetic_training_datasets.py",
        "seed": SEED,
        "source_workbook": RAW_WORKBOOK.name,
        "source_status": SYNTHETIC_STATUS,
        "record_counts": {
            "station_crosswalk": len(crosswalk_rows),
            "passenger_volume_monthly": len(forecast_rows),
            "station_master": len(cluster_rows),
            "station_month_features": len(regression_rows),
            "location_opportunity_training": len(classification_rows),
            "station_insight_context": len(llm_context_rows),
            "smart_query_evaluation": len(llm_eval_rows),
        },
        "limitations": [
            "All generated outputs are synthetic and training-only.",
            "Regression target and classification label are formula-generated, not observed outcomes.",
            "LLM context supports prompt evaluation, not local LLM model training.",
        ],
    }
    manifest_path = ROOT / "ml" / "shared" / "manifests" / "synthetic_dataset_manifest.json"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest["record_counts"], ensure_ascii=False))


if __name__ == "__main__":
    main()
